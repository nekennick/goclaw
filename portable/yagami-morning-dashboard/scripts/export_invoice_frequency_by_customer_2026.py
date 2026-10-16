import csv, datetime as dt, gzip, json, time, urllib.parse, urllib.request
from pathlib import Path
from collections import defaultdict
from statistics import mean
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT=Path.cwd(); DATA=ROOT/'data/kiotviet'; TZ=dt.timezone(dt.timedelta(hours=7))
NOW=dt.datetime.now(TZ).replace(tzinfo=None); START=dt.datetime(NOW.year,1,1)
OUT=ROOT/'outputs'/('tan-suat-lay-hang-hoa-don-hoan-thanh-%s-den-%s.xlsx'%(NOW.year,NOW.strftime('%Y%m%d-%H%M')))
WAREHOUSE_NAME={'CAO_LANH':'Tổng kho YAGAMI','BINH_DUONG':'KHO YAGAMI BÌNH DƯƠNG'}

def env():
 d={}
 for line in (ROOT/'config/.env.kiotviet').read_text(encoding='utf-8').splitlines():
  line=line.strip()
  if line and not line.startswith('#') and '=' in line:
   k,v=line.split('=',1); d[k.strip()]=v.strip().strip('"').strip("'")
 return d

def hj(req,to=90):
 with urllib.request.urlopen(req,timeout=to) as r: return json.loads(r.read().decode('utf-8'))

def token(e):
 data=urllib.parse.urlencode({'scopes':e.get('KIOTVIET_SCOPE','PublicApi.Access'),'grant_type':'client_credentials','client_id':e['KIOTVIET_CLIENT_ID'],'client_secret':e['KIOTVIET_CLIENT_SECRET']}).encode()
 return hj(urllib.request.Request(e.get('KIOTVIET_AUTH_URL','https://id.kiotviet.vn/connect/token'),data=data,headers={'Content-Type':'application/x-www-form-urlencoded'},method='POST'),30)['access_token']

def api(e,t,path,params):
 url=e.get('KIOTVIET_API_URL','https://public.kiotapi.com').rstrip()+path
 if params: url+='?'+urllib.parse.urlencode(params)
 return hj(urllib.request.Request(url,headers={'Authorization':'Bearer '+t,'Retailer':e['KIOTVIET_RETAILER']}),90)

def parse(s):
 try: return dt.datetime.fromisoformat(str(s or '')[:19])
 except: return None

def fetch_all(e,t,path,params,max_pages=200):
 out=[]; cur=0; ps=int(params.get('pageSize',100)); total=None; pages=0
 while True:
  q=dict(params); q['currentItem']=cur
  obj=api(e,t,path,q); arr=obj.get('data',[]) if isinstance(obj,dict) else []
  if total is None and isinstance(obj,dict): total=obj.get('total')
  if not arr: break
  out.extend(arr); cur+=len(arr); pages+=1
  if total is not None and cur>=int(total): break
  if len(arr)<ps or pages>=max_pages: break
  time.sleep(0.04)
 return out,total,pages

def is_done_hd(i):
 code=str(i.get('code') or '').upper(); st=str(i.get('statusValue') or '').strip().upper()
 return st=='HOÀN THÀNH' and (code.startswith('HD') or code.startswith('HĐ'))

def money(x):
 try: return float(x or 0)
 except: return 0.0

def fmt(d): return d.strftime('%d/%m') if d else ''
def classify(n, months, avg):
 if n>=80 or (avg and avg<=5): return 'Lấy rất thường xuyên'
 if n>=30 or (avg and avg<=10): return 'Lấy hằng tuần/ổn định'
 if n>=10 or months>=6: return 'Lấy định kỳ'
 if n>=3: return 'Lấy thưa'
 if n>=1: return 'Rất ít đơn'
 return 'Chưa có HĐ hoàn thành trong năm'

def write_raw(name,items,stamp,synced):
 d=DATA/'raw'/dt.datetime.now(TZ).strftime('%Y-%m-%d'); d.mkdir(parents=True,exist_ok=True)
 with gzip.open(d/f'{name}-{stamp}.jsonl.gz','wt',encoding='utf-8') as f:
  for it in items: f.write(json.dumps({'source':'KiotViet','endpoint':name,'syncId':stamp,'syncedAt':synced,'data':it},ensure_ascii=False)+'\n')

def main():
 e=env(); t=token(e); stamp=dt.datetime.now(dt.timezone.utc).strftime('%Y%m%dT%H%M%SZ'); synced=dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat()
 customers,_,_=fetch_all(e,t,'/customers',{'pageSize':100},max_pages=20)
 kb_by_code={}; kb_path=ROOT/'data/yagami-knowledge-base/yagami-branch-warehouse-map.json'
 if kb_path.exists():
  kb=json.load(open(kb_path,encoding='utf-8'))
  for b in kb.get('branches',[]):
   if b.get('customerCode'): kb_by_code.setdefault(b['customerCode'],b)
 all_invoices=[]; rows=[]; details=[]
 for idx,c in enumerate(customers,1):
  cc=c.get('code')
  if not cc: continue
  params={'pageSize':100,'customerCode':cc,'fromPurchaseDate':START.strftime('%Y-%m-%dT%H:%M:%S'),'toPurchaseDate':NOW.strftime('%Y-%m-%dT%H:%M:%S'),'orderBy':'purchaseDate','orderDirection':'Desc'}
  invs,_,_=fetch_all(e,t,'/invoices',params,max_pages=20)
  invs=[i for i in invs if parse(i.get('purchaseDate')) and START<=parse(i.get('purchaseDate'))<=NOW]
  done=sorted([i for i in invs if is_done_hd(i)], key=lambda x: parse(x.get('purchaseDate')) or dt.datetime.min)
  all_invoices += done
  md={m:[] for m in range(1,13)}; mc={m:[] for m in range(1,13)}; mv={m:0.0 for m in range(1,13)}
  for inv in done:
   p=parse(inv.get('purchaseDate')); md[p.month].append(p); mc[p.month].append(inv.get('code') or ''); mv[p.month]+=money(inv.get('total'))
   details.append([cc,c.get('name') or inv.get('customerName'),inv.get('code'),inv.get('orderCode') or '',p.strftime('%d/%m/%Y %H:%M'),p.month,inv.get('branchName'),inv.get('statusValue'),inv.get('saleChannelId') or inv.get('SaleChannelName') or '',money(inv.get('total'))])
  dates=[parse(x.get('purchaseDate')) for x in done if parse(x.get('purchaseDate'))]
  gaps=[(b-a).days for a,b in zip(dates,dates[1:])]; avg=round(mean(gaps),1) if gaps else None
  months=sum(1 for m in md if md[m]); kb=kb_by_code.get(cc,{})
  row={'Mã KH':cc,'Tên chi nhánh/khách hàng':c.get('name') or '','Kho phụ trách':WAREHOUSE_NAME.get(kb.get('warehouse'),kb.get('warehouse') or ''),'Ngày lịch':kb.get('day') or '','Loại tuyến':kb.get('routeType') or '','Tổng HĐ hoàn thành':len(done),'Số tháng có lấy hàng':months,'Ngày HĐ đầu tiên':dates[0].strftime('%d/%m/%Y') if dates else '','Ngày HĐ gần nhất':dates[-1].strftime('%d/%m/%Y') if dates else '','Khoảng cách TB giữa HĐ (ngày)':avg if avg is not None else '','Nhận định tần suất':classify(len(done),months,avg),'Tổng giá trị HĐ hoàn thành':sum(mv.values())}
  for m in range(1,13):
   row[f'Tháng {m} - Số HĐ']=len(md[m]); row[f'Tháng {m} - Ngày HĐ']=', '.join(fmt(x) for x in md[m]); row[f'Tháng {m} - Mã HĐ']=', '.join(mc[m])
  rows.append(row)
  if idx%25==0: print('processed',idx,'/',len(customers),flush=True)
 write_raw('invoices-complete-frequency-by-customer-current-year', all_invoices, stamp, synced)
 wb=Workbook(); ws=wb.active; ws.title='Tan suat theo thang'
 rows.sort(key=lambda r:(r['Kho phụ trách'] or 'ZZZ',r['Ngày lịch'] or 'ZZZ',r['Tên chi nhánh/khách hàng'] or ''))
 headers=list(rows[0].keys()) if rows else []; ws.append(headers)
 for r in rows: ws.append([r.get(h,'') for h in headers])
 ws2=wb.create_sheet('Chi tiet hoa don'); ws2.append(['Mã KH','Tên chi nhánh/khách hàng','Mã HĐ','Mã đơn gốc','Ngày HĐ','Tháng','Kho xử lý','Trạng thái','Kênh bán/ID','Giá trị'])
 for d in sorted(details,key=lambda x:(x[1] or '',x[4])): ws2.append(d)
 ws3=wb.create_sheet('Summary')
 summary=[['Mốc dữ liệu','Từ %s đến %s'%(START.strftime('%d/%m/%Y'),NOW.strftime('%d/%m/%Y %H:%M'))],['Dữ liệu đồng bộ lúc',dt.datetime.fromisoformat(synced).astimezone(TZ).strftime('%d/%m/%Y %H:%M')],['Tiêu chí tính','Chỉ hóa đơn trạng thái Hoàn thành và mã bắt đầu bằng HD/HĐ'],['Tổng khách hàng hiện tại từ KiotViet',len(customers)],['Tổng HĐ hoàn thành dùng tính tần suất',len(all_invoices)]]
 for r in summary: ws3.append(r)
 ws3.append([]); ws3.append(['Nhóm tần suất','Số khách hàng']); cnt=defaultdict(int)
 for r in rows: cnt[r['Nhận định tần suất']]+=1
 for k in ['Lấy rất thường xuyên','Lấy hằng tuần/ổn định','Lấy định kỳ','Lấy thưa','Rất ít đơn','Chưa có HĐ hoàn thành trong năm']: ws3.append([k,cnt[k]])
 header_fill=PatternFill('solid',fgColor='1F4E78'); header_font=Font(name='Arial',bold=True,color='FFFFFF'); normal_font=Font(name='Arial',size=10); thin=Side(style='thin',color='D9E2F3')
 for sh in wb.worksheets:
  for row in sh.iter_rows():
   for cell in row:
    cell.font=normal_font; cell.alignment=Alignment(vertical='top',wrap_text=True); cell.border=Border(left=thin,right=thin,top=thin,bottom=thin)
  for cell in sh[1]: cell.font=header_font; cell.fill=header_fill; cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
  sh.freeze_panes='A2'; sh.auto_filter.ref=sh.dimensions
  for col in sh.columns:
   sh.column_dimensions[get_column_letter(col[0].column)].width=min(max(max(len(str(c.value or '')) for c in col)+2,10),55)
 OUT.parent.mkdir(exist_ok=True); wb.save(OUT)
 print('DONE')
 print('Moc:',START.strftime('%d/%m/%Y'),'den',NOW.strftime('%d/%m/%Y %H:%M'))
 print('Khach hang:',len(customers))
 print('HD hoan thanh:',len(all_invoices))
 print('File:',OUT)
if __name__=='__main__': main()
