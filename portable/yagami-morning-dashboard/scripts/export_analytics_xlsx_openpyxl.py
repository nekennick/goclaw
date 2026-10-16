from pathlib import Path
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

root = Path(__file__).resolve().parents[1]
data = json.loads((root / 'data/analytics/analytics.json').read_text(encoding='utf-8'))['stoppedProducts']
out = root / 'outputs/bao-cao-mat-hang-ngung-mua-hoa-don-2026.xlsx'
wb = Workbook()
ws = wb.active
ws.title = 'Kết quả phân tích'
headers = ['Khách hàng / chi nhánh','Mặt hàng','Đơn từ đầu năm','Hóa đơn gần nhất','Ngày mua gần nhất','Đã cách (ngày)']
ws.append(headers)
for x in data.get('items', []):
    ws.append([x.get('branch'), x.get('product'), x.get('baselineOrders'), x.get('latestInvoice'), x.get('latestPurchaseDate'), x.get('daysSincePurchase')])

alias = wb.create_sheet('Mã hàng đã gom')
alias.append(['Mã đang hoạt động','Mã cũ','Tên sản phẩm'])
for item in data.get('productAliases', []):
    for old in item.get('retiredCodes', []): alias.append([item.get('activeCode'), old, item.get('name')])

excluded = wb.create_sheet('Mã hàng bị loại')
excluded.append(['Mã bị loại','Lý do'])
for code in data.get('ignoredInactiveCodes', []): excluded.append([code, 'Mã sản phẩm inactive'])
excluded.append(['{DEL}', 'Mã đã xóa / đánh dấu xóa'])

info = wb.create_sheet('Thông tin báo cáo')
for row in [
    ['BÁO CÁO MẶT HÀNG KHÁCH HÀNG NGỪNG MUA'],
    ['Nguồn', data.get('source')],
    ['Dữ liệu từ', data.get('dataFrom')],
    ['Dữ liệu đến', data.get('dataTo')],
    ['Giai đoạn nền', f"{data['baseline']['from']} đến {data['baseline']['to']}"],
    ['30 ngày gần nhất', f"{data['recent']['from']} đến {data['recent']['to']}"],
    ['Số hóa đơn đã đọc', data.get('rows')],
    ['Số trường hợp cần kiểm tra', data.get('stoppedCount')],
    ['Quy tắc', 'Tối thiểu 2 hóa đơn từ đầu năm và không có hóa đơn trong 30 ngày gần nhất'],
    ['Chuẩn hóa', 'Gom theo customerId; mã cũ có mã hoạt động thay thế được gom về mã mới; mã inactive và {DEL} bị loại'],
    ['Loại mã kết thúc bằng -', 'Không thống kê các mã như MDT-'],
]: info.append(row)

header_fill = PatternFill('solid', fgColor='1F4E78')
for sheet in wb.worksheets:
    sheet.freeze_panes = 'A2'
    if sheet.max_row > 1: sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    for row in sheet.iter_rows():
        for cell in row: cell.alignment = Alignment(wrap_text=True, vertical='top')
    for column in sheet.columns:
        letter = get_column_letter(column[0].column)
        longest = max((len(str(cell.value or '')) for cell in column), default=10)
        sheet.column_dimensions[letter].width = min(max(longest + 2, 12), 60)
out.parent.mkdir(parents=True, exist_ok=True)
wb.save(out)
print(json.dumps({'ok': True, 'file': str(out), 'rows': len(data.get('items', []))}, ensure_ascii=False))
