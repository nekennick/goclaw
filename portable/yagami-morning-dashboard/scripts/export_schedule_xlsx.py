from pathlib import Path
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

root = Path(__file__).resolve().parents[1]
snapshot = json.loads((root / 'data/analytics/analytics.json').read_text(encoding='utf-8'))
schedule = snapshot['schedule']
out = root / 'outputs/bao-cao-chi-nhanh-khong-dat-dung-lich-2026.xlsx'
wb = Workbook()

ranking = wb.active
ranking.title = 'Xếp hạng chi nhánh'
ranking.append(['Chi nhánh', 'Số ngày chưa đặt', 'Mức độ'])
for index, item in enumerate(schedule.get('branchRanking', []), 1):
    ranking.append([item['branch'], item['days'], 'Ưu tiên' if index <= 3 else 'Theo dõi'])

daily = wb.create_sheet('Theo ngày')
daily.append(['Ngày báo cáo', 'Kho phụ trách', 'Chi nhánh chưa đặt', 'Số chi nhánh'])
for item in schedule.get('daily', []):
    for warehouse, branches in item.get('warehouses', {}).items():
        daily.append([item['date'], warehouse, ', '.join(branches) if branches else 'Không có', len(branches)])

info = wb.create_sheet('Thông tin báo cáo')
for row in [
    ['BÁO CÁO CHI NHÁNH KHÔNG ĐẶT ĐÚNG LỊCH'],
    ['Nguồn', schedule.get('source')],
    ['Mốc bắt đầu tin cậy', schedule.get('firstReliableDate')],
    ['Mốc dữ liệu cuối', schedule.get('lastDate')],
    ['Số ngày hợp lệ', schedule.get('validDays')],
    ['Lần chạy bị loại', schedule.get('ignoredRuns')],
    ['Ngày loại khỏi phân tích', ', '.join(schedule.get('excludedDates', []))],
    ['Quy tắc', 'Chỉ tính lần cron thành công có rõ mục Chi nhánh chưa đặt; ngày chạy trùng chỉ giữ lần mới nhất'],
    ['Lưu ý', 'Đây là tình trạng chưa đặt tại mốc báo cáo sáng, không kết luận chi nhánh không đặt cả ngày'],
]: info.append(row)

fill = PatternFill('solid', fgColor='1F4E78')
for sheet in wb.worksheets:
    sheet.freeze_panes = 'A2'
    if sheet.max_row > 1: sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    for row in sheet.iter_rows():
        for cell in row: cell.alignment = Alignment(wrap_text=True, vertical='top')
    for column in sheet.columns:
        letter = get_column_letter(column[0].column)
        longest = max((len(str(cell.value or '')) for cell in column), default=10)
        sheet.column_dimensions[letter].width = min(max(longest + 2, 12), 70)
out.parent.mkdir(parents=True, exist_ok=True)
wb.save(out)
print(json.dumps({'ok': True, 'file': str(out), 'rankingRows': len(schedule.get('branchRanking', [])), 'dailyRows': daily.max_row - 1}, ensure_ascii=False))
