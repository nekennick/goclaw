import csv, datetime as dt, gzip, json, time, urllib.parse, urllib.request
from pathlib import Path
from collections import defaultdict
from statistics import mean
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT=Path.cwd()
DATA=ROOT/'data/kiotviet'
TZ=dt.timezone(dt.timedelta(hours=7))
NOW=dt.datetime.now(TZ).replace(tzinfo=None)
START=dt.datetime(NOW.year,1,1,0,0,0)
OUT=ROOT/'outputs'/('tan-suat-dat-hang-hoa-don-hoan-thanh-%s-den-%s.xlsx' % (NOW.year, NOW.strftime('%Y%m%d-%H%M')))
WAREHOUSE_NAME={'CAO_LANH':'Tổng kho YAGAMI','BINH_DUONG':'KHO YAGAMI BÌNH DƯƠNG'}

def load_env():
    d={}
    for line in (ROOT/'config/.env.kiotviet').read_text(encoding='utf-8').splitlines():
        line=line.strip()
        if line and not line.startswith('#') and '=' in line:
            k,v=line.split('=',1); d[k.strip()]=v.strip().strip('"').strip("'")
    return d

def http_json(req, timeout=90):
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return json.loads(r.read().decode('utf-8'))

def get_token(env):
    data=urllib.parse.urlencode({'scopes':env.get('KIOTVIET_SCOPE','PublicApi.Access'),'grant_type':'client_credentials','client_id':env['KIOTVIET_CLIENT_ID'],'client_secret':env['KIOTVIET_CLIENT_SECRET']}).encode()
    req=urllib.request.Request(env.get('KIOTVIET_AUTH_URL','https://id.kiotviet.vn/connect/token'), data=data, headers={'Content-Type':'application/x-www-form-urlencoded'}, method='POST')
    return http_json(req,30)['access_token']

def api(env, token, path, params):
    url=env.get('KIOTVIET_API_URL','https://public.kiotapi.com').rstrip()+path
    if params: url+='?'+urllib.parse.urlencode(params)
    req=urllib.request.Request(url, headers={'Authorization':'Bearer '+token,'Retailer':env['KIOTVIET_RETAILER']})
    return http_json(req,90)

def parse_date(s):
    if not s: return None
    try: return dt.datetime.fromisoformat(str(s)[:19])
    except Exception: return None

def fetch_paged(env, token, path, params, stop_before=None, max_pages=3000):
    out=[]; cur=0; ps=int(params.get('pageSize',100)); pages=0; total=None
    while True:
        q=dict(params); q['currentItem']=cur
        obj=api(env, token, path, q)
        arr=obj.get('data',[]) if isinstance(obj,dict) else []
        if total is None and isinstance(obj,dict): total=obj.get('total')
        if not arr: break
        dates=[]; keep=[]
        for it in arr:
            p=parse_date(it.get('purchaseDate') or it.get('createdDate'))
            if p: dates.append(p)
            if not stop_before or (p and p>=stop_before): keep.append(it)
        out.extend(keep); pages+=1; cur+=len(arr)
        if stop_before and dates and max(dates) < stop_before: break
        if total is not None and cur>=int(total): break
        if len(arr)<ps or pages>=max_pages: break
        time.sleep(0.05)
    return out,total,pages

def flat(x):
    return {k:(json.dumps(v,ensure_ascii=False) if isinstance(v,(list,dict)) else v) for k,v in x.items()}

def write_csv(path, rows):
    path.parent.mkdir(parents=True,exist_ok=True); fields=[]
    for r in rows:
        for k in r:
            if k not in fields: fields.append(k)
    with path.open('w',encoding='utf-8',newline='') as f:
        w=csv.DictWriter(f,fieldnames=fields,extrasaction='ignore'); w.writeheader(); w.writerows(rows)

def raw_save(name, items, stamp, synced):
    d=DATA/'raw'/dt.datetime.now(TZ).strftime('%Y-%m-%d'); d.mkdir(parents=True,exist_ok=True)
    p=d/f'{name}-{stamp}.jsonl.gz'
    with gzip.open(p,'wt',encoding='utf-8') as f:
        for it in items: f.write(json.dumps({'source':'KiotViet','endpoint':name,'syncId':stamp,'syncedAt':synced,'data':it},ensure_ascii=False)+'\n')
    return p

def fmt_date(d): return d.strftime('%d/%m') if d else ''
def money(x):
    try: return float(x or 0)
    except Exception: return 0.0

def is_complete_invoice(i):
    code=str(i.get('code') or '').upper()
    status=str(i.get('statusValue') or '').strip().upper()
    return status=='HOÀN THÀNH' and (code.startswith('HD') or code.startswith('HĐ'))

def classify(total_docs, months_active, avg_gap):
    if total_docs >= 80 or (avg_gap and avg_gap <= 5): return 'Lấy rất thường xuyên'
    if total_docs >= 30 or (avg_gap and avg_gap <= 10): return 'Lấy hằng tuần/ổn định'
    if total_docs >= 10 or months_active >= 6: return 'Lấy định kỳ'
    if total_docs >= 3: return 'Lấy thưa'
    if total_docs >= 1: return 'Rất ít đơn'
    return 'Chưa có HĐ hoàn thành trong năm'

def main():
    env=load_env(); token=get_token(env)
    stamp=dt.datetime.now(dt.timezone.utc).strftime('%Y%m%dT%H%M%SZ')
    synced=dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat()
    customers,total_c,pages_c=fetch_paged(env, token, '/customers', {'pageSize':100})
    invoices,total_i,pages_i=fetch_paged(env, token, '/invoices', {'pageSize':100,'orderBy':'purchaseDate','orderDirection':'Desc','fromPurchaseDate':START.strftime('%Y-%m-%dT%H:%M:%S'),'toPurchaseDate':NOW.strftime('%Y-%m-%dT%H:%M:%S')}, stop_before=START)
    invoices=[i for i in invoices if parse_date(i.get('purchaseDate')) and START <= parse_date(i.get('purchaseDate')) <= NOW]
    complete=[i for i in invoices if is_complete_invoice(i)]
    raw_save('customers-invoice-frequency-report', customers, stamp, synced)
    raw_save('invoices-frequency-report-current-year', invoices, stamp, synced)
    write_csv(DATA/'normalized/invoices-frequency-report-current-year.csv',[flat(x) for x in invoices])

    kb_by_code={}; kb_path=ROOT/'data/yagami-knowledge-base/yagami-branch-warehouse-map.json'
    if kb_path.exists():
        kb=json.load(open(kb_path,encoding='utf-8'))
        for b in kb.get('branches',[]):
            if b.get('customerCode'): kb_by_code.setdefault(b['customerCode'],b)
    cust_by_code={}
    for c in customers:
        if c.get('code'): cust_by_code[c.get('code')]=c
    for i in complete:
        cc=i.get('customerCode')
        if cc and cc not in cust_by_code: cust_by_code[cc]={'code':cc,'name':i.get('customerName') or cc,'_missingInCustomerList':True}
    inv_by_code=defaultdict(list)
    for i in complete:
        if i.get('customerCode'): inv_by_code[i.get('customerCode')].append(i)
    rows=[]; details=[]
    for cc,c in cust_by_code.items():
        arr=sorted(inv_by_code.get(cc,[]), key=lambda x: parse_date(x.get('purchaseDate')) or dt.datetime.min)
        month_dates={m:[] for m in range(1,13)}; month_codes={m:[] for m in range(1,13)}; month_value={m:0.0 for m in range(1,13)}
        for inv in arr:
            p=parse_date(inv.get('purchaseDate'))
            if not p: continue
            month_dates[p.month].append(p); month_codes[p.month].append(inv.get('code') or ''); month_value[p.month]+=money(inv.get('total'))
            details.append({'customerCode':cc,'customerName':c.get('name') or inv.get('customerName'),'invoiceCode':inv.get('code'),'orderCode':inv.get('orderCode') or '','purchaseDate':p.strftime('%d/%m/%Y %H:%M'),'month':p.month,'branchName':inv.get('branchName'),'status':inv.get('statusValue'),'saleChannel':inv.get('SaleChannelName') or inv.get('saleChannelId') or inv.get('saleChannelId') or '','total':money(inv.get('total'))})
        dates=[parse_date(x.get('purchaseDate')) for x in arr if parse_date(x.get('purchaseDate'))]
        gaps=[(b-a).days for a,b in zip(dates, dates[1:])]
        total_docs=len(arr); months_active=sum(1 for m in range(1,13) if month_dates[m]); avg_gap=round(mean(gaps),1) if gaps else None
        kb=kb_by_code.get(cc,{})
        row={'Mã KH':cc,'Tên chi nhánh/khách hàng':c.get('name') or '','Kho phụ trách':WAREHOUSE_NAME.get(kb.get('warehouse'), kb.get('warehouse') or ''),'Ngày lịch':kb.get('day') or '','Loại tuyến':kb.get('routeType') or '','Tổng HĐ hoàn thành':total_docs,'Số tháng có lấy hàng':months_active,'Ngày HĐ đầu tiên':dates[0].strftime('%d/%m/%Y') if dates else '','Ngày HĐ gần nhất':dates[-1].strftime('%d/%m/%Y') if dates else '','Khoảng cách TB giữa HĐ (ngày)':avg_gap if avg_gap is not None else '','Nhận định tần suất':classify(total_docs, months_active, avg_gap),'Tổng giá trị HĐ hoàn thành':sum(month_value.values())}
        for m in range(1,13):
            row[f'Tháng {m} - Số HĐ']=len(month_dates[m]); row[f'Tháng {m} - Ngày HĐ']=', '.join(fmt_date(x) for x in month_dates[m]); row[f'Tháng {m} - Mã HĐ']=', '.join(month_codes[m])
        rows.append(row)
    rows.sort(key=lambda r:(r['Kho phụ trách'] or 'ZZZ', r['Ngày lịch'] or 'ZZZ', r['Tên chi nhánh/khách hàng'] or ''))

    wb=Workbook(); ws=wb.active; ws.title='Tan suat theo thang'
    headers=list(rows[0].keys()) if rows else []; ws.append(headers)
    for r in rows: ws.append([r.get(h,'') for h in headers])
    ws2=wb.create_sheet('Chi tiet hoa don')
    dh=['Mã KH','Tên chi nhánh/khách hàng','Mã HĐ','Mã đơn gốc','Ngày HĐ','Tháng','Kho xử lý','Trạng thái','Kênh bán/ID','Giá trị']
    ws2.append(dh)
    for d in sorted(details,key=lambda x:(x['customerName'] or '', x['purchaseDate'])):
        ws2.append([d['customerCode'],d['customerName'],d['invoiceCode'],d['orderCode'],d['purchaseDate'],d['month'],d['branchName'],d['status'],d['saleChannel'],d['total']])
    ws3=wb.create_sheet('Summary')
    summary=[['Mốc dữ liệu', 'Từ %s đến %s' % (START.strftime('%d/%m/%Y'), NOW.strftime('%d/%m/%Y %H:%M'))],['Dữ liệu đồng bộ lúc', dt.datetime.fromisoformat(synced).astimezone(TZ).strftime('%d/%m/%Y %H:%M')],['Tiêu chí tính', 'Chỉ hóa đơn có trạng thái Hoàn thành và mã bắt đầu bằng HD/HĐ'],['Tổng khách hàng hiện tại từ KiotViet', len(customers)],['Tổng khách hàng/chi nhánh trong file', len(rows)],['Tổng HĐ trong kỳ đã tải', len(invoices)],['Tổng HĐ hoàn thành dùng tính tần suất', len(complete)],['Số page invoices đã đọc', pages_i],['Ghi chú', 'Mỗi tháng gồm số HĐ, ngày HĐ và mã HĐ; ngày hiển thị dạng dd/mm.']]
    for x in summary: ws3.append(x)
    ws3.append([]); ws3.append(['Nhóm tần suất','Số khách hàng'])
    cnt=defaultdict(int)
    for r in rows: cnt[r['Nhận định tần suất']]+=1
    for k in ['Lấy rất thường xuyên','Lấy hằng tuần/ổn định','Lấy định kỳ','Lấy thưa','Rất ít đơn','Chưa có HĐ hoàn thành trong năm']: ws3.append([k,cnt.get(k,0)])
    header_fill=PatternFill('solid',fgColor='1F4E78'); header_font=Font(name='Arial',bold=True,color='FFFFFF'); normal_font=Font(name='Arial',size=10); thin=Side(style='thin',color='D9E2F3')
    for sh in wb.worksheets:
        for row in sh.iter_rows():
            for cell in row:
                cell.font=normal_font; cell.alignment=Alignment(vertical='top',wrap_text=True); cell.border=Border(left=thin,right=thin,top=thin,bottom=thin)
        for cell in sh[1]: cell.font=header_font; cell.fill=header_fill; cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        sh.freeze_panes='A2'; sh.auto_filter.ref=sh.dimensions
        for col in sh.columns:
            maxlen=max(len(str(c.value or '')) for c in col); sh.column_dimensions[get_column_letter(col[0].column)].width=min(max(maxlen+2,10),55)
    OUT.parent.mkdir(exist_ok=True); wb.save(OUT)
    print('Đã xuất tần suất lấy hàng theo HĐ hoàn thành từ %s đến %s' % (START.strftime('%d/%m/%Y'), NOW.strftime('%d/%m/%Y %H:%M')))
    print('Tổng khách hàng KiotViet:', len(customers))
    print('Tổng khách hàng/chi nhánh trong file:', len(rows))
    print('Tổng HĐ trong kỳ đã tải:', len(invoices))
    print('Tổng HĐ hoàn thành dùng tính:', len(complete))
    print('File:', OUT)
if __name__=='__main__': main()
